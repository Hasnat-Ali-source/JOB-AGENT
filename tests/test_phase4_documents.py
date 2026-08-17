#!/usr/bin/env python3
"""
Tests for Phase 4: Document Generation.

Covers the Phase 4 acceptance criteria:
- A master resume (DOCX or PDF) is uploaded, parsed, and stored
- Given a job posting, a tailored resume PDF variant is generated
- The PDF is linked to the version record and downloadable for review

Plus the safety property the phase depends on: a tailored document must never
assert anything the master document doesn't support.
"""

from pathlib import Path

import pytest
import pytest_asyncio
from sqlalchemy.pool import StaticPool
from sqlmodel import Session, SQLModel, create_engine

from job_agent.models.database import (
    AuditAction,
    AuditLog,
    DocumentFormat,
    DocumentType,
    DocumentVersion,
    Job,
    MasterDocument,
)
from job_agent.services.document_builder import DocumentBuilder
from job_agent.services.document_parser import DocumentParser, parse_document
from job_agent.services.fabrication_check import verify_no_fabrication
from job_agent.services.pdf_renderer import (
    PdfRenderer,
    select_engine,
    weasyprint_available,
)
from job_agent.services.tailoring import TailoringService

MASTER_RESUME = """Alex Rivera
alex.rivera@example.com | +1 555 0100 | San Francisco, CA

Summary
Backend engineer with 8 years building distributed systems.

Experience
Staff Engineer, Northwind Systems (2021 - Present)
- Led migration of the billing platform to PostgreSQL, cutting query latency by 35%
- Designed an event pipeline processing 12M events per day
- Mentored 4 engineers across two teams

Backend Engineer, Lumen Data (2018 - 2021)
- Built REST services in Python and Django
- Introduced automated testing, raising coverage to 82%

Education
BSc Computer Science, University of Washington (2018)

Skills
Python, PostgreSQL, AWS, Kafka, Docker, Kubernetes
"""

MASTER_COVER_LETTER = """Alex Rivera
alex.rivera@example.com

Dear Hiring Manager,

I am a backend engineer with 8 years of experience building distributed systems
in Python and PostgreSQL. At Northwind Systems I led the billing platform
migration and designed an event pipeline processing 12M events per day.

Sincerely,
Alex Rivera
"""


# ============================================================================
# Fixtures
# ============================================================================

@pytest.fixture
def session():
    """In-memory database with every table created."""
    engine = create_engine(
        "sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    SQLModel.metadata.create_all(engine)

    with Session(engine) as db_session:
        yield db_session


@pytest.fixture
def docs_dir(tmp_path) -> Path:
    """Isolated document storage root."""
    return tmp_path / "documents"


@pytest.fixture
def builder(session, docs_dir) -> DocumentBuilder:
    """DocumentBuilder writing into the temporary storage root."""
    return DocumentBuilder(session, documents_dir=docs_dir)


@pytest.fixture
def resume_txt(tmp_path) -> Path:
    """A master resume as plain text."""
    path = tmp_path / "alex_rivera_resume.txt"
    path.write_text(MASTER_RESUME)
    return path


@pytest.fixture
def resume_docx(tmp_path) -> Path:
    """A master resume as a real .docx file."""
    from docx import Document

    document = Document()
    for line in MASTER_RESUME.split("\n"):
        document.add_paragraph(line)

    path = tmp_path / "alex_rivera_resume.docx"
    document.save(str(path))
    return path


@pytest.fixture
def resume_pdf(tmp_path) -> Path:
    """A master resume as a real .pdf file."""
    path = tmp_path / "alex_rivera_resume.pdf"
    PdfRenderer().render(MASTER_RESUME, path, title="Resume")
    return path


@pytest.fixture
def job(session) -> Job:
    """A stored job to tailor against."""
    record = Job(
        platform="generic_ats",
        external_id="REQ-1",
        title="Senior Backend Engineer",
        company="Acme Robotics",
        location="Remote",
        description=(
            "We need a backend engineer experienced with PostgreSQL, Kafka and "
            "event-driven pipelines. You will own billing infrastructure and "
            "mentor other engineers."
        ),
        requirements="Python, PostgreSQL, Kafka, distributed systems",
        salary="USD 180,000 - 210,000",
        apply_method="web_form",
        dedup_hash="hash-1",
    )
    session.add(record)
    session.commit()
    session.refresh(record)
    return record


# ============================================================================
# Parsing
# ============================================================================

class TestDocumentParser:
    """Master documents arrive as .docx, .pdf, .txt or .md."""

    def test_parses_plain_text(self, resume_txt):
        text, sections, doc_format = parse_document(resume_txt)

        assert doc_format == DocumentFormat.TXT
        assert "Alex Rivera" in text
        assert "Northwind Systems" in text

    def test_parses_docx(self, resume_docx):
        text, sections, doc_format = parse_document(resume_docx)

        assert doc_format == DocumentFormat.DOCX
        assert "Alex Rivera" in text
        assert "PostgreSQL" in text

    def test_parses_pdf(self, resume_pdf):
        """Acceptance: a PDF master is parsed, not just accepted."""
        text, sections, doc_format = parse_document(resume_pdf)

        assert doc_format == DocumentFormat.PDF
        assert "Alex Rivera" in text
        assert "Northwind" in text

    def test_a_pdf_is_read_in_reading_order(self, resume_pdf):
        """
        A PDF draws text in whatever order it likes.

        On a designed resume the default extraction returned the contact line,
        the first sentence of the summary and one job title *after* the
        hobbies — so the summary began mid-sentence and a run of bullets sat
        under a heading with no employer above it. Everything downstream
        inherited that: the tailored resume, the checks, the rendered PDF.
        """
        text, _sections, _fmt = parse_document(resume_pdf)

        lines = [line for line in text.splitlines() if line.strip()]
        position = {line: index for index, line in enumerate(lines)}

        name = next(i for line, i in position.items() if "Alex Rivera" in line)
        contact = next(i for line, i in position.items() if "@example.com" in line)
        education = next(
            i for line, i in position.items() if line.strip().lower() == "education"
        )

        assert name < contact < education, f"Out of reading order: {lines[:6]}"

    def test_lines_broken_by_the_page_width_are_rejoined(self, tmp_path):
        """
        A wrapped sentence is one line, not two paragraphs.

        Rendered back at a different width, each fragment became its own
        paragraph — a four-line summary printed as four stubs, and a bullet
        whose last word sat alone underneath it.
        """
        path = tmp_path / "wrapped.txt"
        path.write_text(
            "Alex Rivera\n"
            "alex@example.com\n"
            "\n"
            "Summary\n"
            "Backend engineer with eight years of experience building\n"
            "distributed systems and mentoring other engineers.\n"
            "\n"
            "Experience\n"
            "Staff Engineer, Northwind Systems (2021 - Present)\n"
            "- Led the billing migration to PostgreSQL, cutting query\n"
            "latency by 35%.\n"
        )

        text, _sections, _fmt = parse_document(path)

        assert "building distributed systems" in text
        assert "cutting query latency by 35%." in text
        assert "\nlatency" not in text

    @pytest.mark.parametrize("previous,line", [
        ("Staff Engineer, Northwind Systems (2021 - Present)", "Backend Engineer, Lumen Data"),
        ("Skilled in communication.", "and performance management"),
        ("+60 10 421 5890", "linkedin.com/in/alex-rivera"),
        ("Experience", "staff engineer"),
    ])
    def test_separate_lines_are_left_separate(self, previous, line):
        """Joining too eagerly would run a resume's facts together."""
        assert not DocumentParser._is_wrapped(previous, line)

    def test_detects_sections(self, resume_txt):
        _text, sections, _fmt = parse_document(resume_txt)

        assert "header" in sections  # name and contact details
        assert "summary" in sections
        assert "experience" in sections
        assert "education" in sections
        assert "skills" in sections
        assert "Northwind Systems" in sections["experience"]

    def test_heading_detection(self):
        assert DocumentParser.is_heading("Experience") == "experience"
        assert DocumentParser.is_heading("WORK EXPERIENCE") == "experience"
        assert DocumentParser.is_heading("Technical Skills") == "skills"
        assert DocumentParser.is_heading("Professional Summary") == "summary"

    def test_non_headings_rejected(self):
        assert DocumentParser.is_heading("- Built REST services in Python") is None
        assert DocumentParser.is_heading("I led a team of four engineers.") is None
        assert DocumentParser.is_heading("") is None

    def test_unsupported_format_rejected(self, tmp_path):
        path = tmp_path / "resume.pages"
        path.write_text("x")

        with pytest.raises(ValueError, match="Unsupported document format"):
            parse_document(path)

    def test_empty_document_rejected(self, tmp_path):
        path = tmp_path / "empty.txt"
        path.write_text("   \n  \n")

        with pytest.raises(ValueError, match="No text could be extracted"):
            parse_document(path)

    def test_missing_file_rejected(self, tmp_path):
        with pytest.raises(FileNotFoundError):
            parse_document(tmp_path / "nope.txt")


# ============================================================================
# PDF rendering
# ============================================================================

class TestPdfRenderer:
    """ReportLab output must be a real, non-trivial PDF."""

    def test_renders_a_pdf_file(self, tmp_path):
        out = tmp_path / "out.pdf"
        PdfRenderer().render(MASTER_RESUME, out, title="Resume")

        assert out.exists()
        assert out.read_bytes().startswith(b"%PDF")
        assert out.stat().st_size > 1000

    def test_creates_missing_directories(self, tmp_path):
        out = tmp_path / "a" / "b" / "c" / "out.pdf"
        PdfRenderer().render("Alex Rivera\n\nSummary\nEngineer.", out)

        assert out.exists()

    def test_empty_text_rejected(self, tmp_path):
        with pytest.raises(ValueError, match="empty document"):
            PdfRenderer().render("   ", tmp_path / "out.pdf")

    def test_markup_characters_do_not_break_rendering(self, tmp_path):
        """A resume mentioning C++ & <framework> must not crash the renderer."""
        out = tmp_path / "escaped.pdf"
        text = "Alex Rivera\n\nSkills\n- C++ & Python <3\n- Uses A&B <tags>"

        PdfRenderer().render(text, out)

        assert out.exists()
        assert out.read_bytes().startswith(b"%PDF")

    def test_round_trips_through_pdf_extraction(self, tmp_path):
        """Rendered text must be extractable — proof it isn't an empty page."""
        out = tmp_path / "round.pdf"
        PdfRenderer().render(MASTER_RESUME, out)

        text, _sections, _fmt = parse_document(out)

        assert "Alex Rivera" in text
        assert "Northwind" in text


ENGINES = ["reportlab"] + (["weasyprint"] if weasyprint_available() else [])


class TestRenderEngines:
    """
    Both engines must produce the same document.

    WeasyPrint gives better typography but needs pango/cairo; ReportLab is the
    pure-Python fallback. Whichever runs, the content has to be identical and
    correct — these run against every engine available on the machine.
    """

    @pytest.mark.parametrize("engine", ENGINES)
    def test_engine_produces_a_valid_pdf(self, engine, tmp_path):
        out = tmp_path / f"{engine}.pdf"
        PdfRenderer(engine=engine).render(MASTER_RESUME, out, title="Resume")

        assert out.read_bytes().startswith(b"%PDF")

    @pytest.mark.parametrize("engine", ENGINES)
    def test_engine_preserves_all_content(self, engine, tmp_path):
        out = tmp_path / f"{engine}.pdf"
        PdfRenderer(engine=engine).render(MASTER_RESUME, out, title="Resume")

        text, _sections, _fmt = parse_document(out)

        for expected in ("Alex Rivera", "alex.rivera@example.com", "Northwind Systems",
                         "Lumen Data", "University of Washington", "Kubernetes"):
            assert expected in text, f"{engine} dropped '{expected}'"

    @pytest.mark.parametrize("engine", ENGINES)
    def test_bullet_lines_survive_rendering(self, engine, tmp_path):
        """
        Bullets must render as glyphs, not as the literal word "bullet".

        ReportLab's ListItem takes a `value` argument that silently prints its
        own text as the marker — a PDF that is structurally valid and visually
        wrong. Checking only for %PDF bytes would not catch it.
        """
        out = tmp_path / f"{engine}_bullets.pdf"
        PdfRenderer(engine=engine).render(MASTER_RESUME, out)

        text, _sections, _fmt = parse_document(out)

        assert "Mentored 4 engineers" in text
        assert "bullet" not in text.lower()

    @pytest.mark.parametrize("engine", ENGINES)
    def test_markup_characters_survive_every_engine(self, engine, tmp_path):
        out = tmp_path / f"{engine}_escape.pdf"
        PdfRenderer(engine=engine).render(
            "Alex Rivera\n\nSkills\n- C++ & Python <3\n- Uses A&B <tags>", out)

        text, _sections, _fmt = parse_document(out)

        assert "C++" in text
        assert "&amp;" not in text  # escaped for the renderer, not for the reader

    def test_explicit_reportlab_is_honoured(self):
        assert select_engine("reportlab") == "reportlab"

    def test_auto_prefers_weasyprint_when_available(self):
        expected = "weasyprint" if weasyprint_available() else "reportlab"
        assert select_engine("auto") == expected

    def test_weasyprint_request_degrades_when_unavailable(self, monkeypatch):
        """Pinning an engine that can't run must fall back, not crash."""
        monkeypatch.setattr(
            "job_agent.services.pdf_renderer.weasyprint_available", lambda: False)

        assert select_engine("weasyprint") == "reportlab"

    def test_isolated_render_produces_a_valid_pdf(self, tmp_path):
        """WeasyPrint in a subprocess must produce the same output."""
        if not weasyprint_available():
            pytest.skip("WeasyPrint cannot render here")

        out = tmp_path / "isolated.pdf"
        renderer = PdfRenderer(engine="weasyprint", isolate=True)
        renderer.render(MASTER_RESUME, out, title="Resume")

        assert renderer.engine == "weasyprint"  # did not silently fall back
        assert out.read_bytes().startswith(b"%PDF")

        text, _sections, _fmt = parse_document(out)
        assert "Alex Rivera" in text

    def test_isolated_render_works_from_any_directory(self, tmp_path, monkeypatch):
        """
        The child must import job_agent without help from the working
        directory — a scheduled run starts wherever launchd puts it.
        """
        if not weasyprint_available():
            pytest.skip("WeasyPrint cannot render here")

        monkeypatch.chdir(tmp_path)

        out = tmp_path / "elsewhere.pdf"
        renderer = PdfRenderer(engine="weasyprint", isolate=True)
        renderer.render(MASTER_RESUME, out)

        assert renderer.engine == "weasyprint"
        assert out.read_bytes().startswith(b"%PDF")

    def test_a_crashing_child_falls_back_instead_of_dying(self, tmp_path, monkeypatch):
        """
        The reason isolation exists.

        A native fault in WeasyPrint's cffi/pango stack kills the whole
        interpreter. In a subprocess it's a negative return code, and the
        document still gets produced.
        """
        class Killed:
            returncode = -11  # SIGSEGV
            stdout = ""
            stderr = ""

        monkeypatch.setattr(
            "job_agent.services.pdf_renderer.subprocess.run",
            lambda *a, **k: Killed(),
        )

        out = tmp_path / "survived.pdf"
        renderer = PdfRenderer(engine="weasyprint", isolate=True)
        renderer.render(MASTER_RESUME, out)

        assert renderer.engine == "reportlab"
        assert out.read_bytes().startswith(b"%PDF")

    def test_a_hanging_child_times_out_and_falls_back(self, tmp_path, monkeypatch):
        import subprocess as sp

        def hang(*args, **kwargs):
            raise sp.TimeoutExpired(cmd="worker", timeout=60)

        monkeypatch.setattr("job_agent.services.pdf_renderer.subprocess.run", hang)

        out = tmp_path / "timeout.pdf"
        renderer = PdfRenderer(engine="weasyprint", isolate=True)
        renderer.render(MASTER_RESUME, out)

        assert renderer.engine == "reportlab"
        assert out.read_bytes().startswith(b"%PDF")

    def test_worker_reports_a_bad_job(self):
        """The worker fails loudly rather than writing nothing silently."""
        import subprocess
        import sys

        completed = subprocess.run(
            [sys.executable, "-m", "job_agent.services._pdf_worker"],
            input="{}", capture_output=True, text=True, timeout=60,
        )

        assert completed.returncode != 0
        assert "html" in completed.stderr

    def test_html_escapes_markup_but_keeps_bold(self):
        from job_agent.services.document_layout import Block

        html_out = PdfRenderer(engine="reportlab").build_html(
            [Block(kind="paragraph", text="A & B <tag> **bold**")], "t")

        assert "&amp;" in html_out
        assert "&lt;tag&gt;" in html_out
        assert "<strong>bold</strong>" in html_out


# ============================================================================
# Fabrication check
# ============================================================================

class TestFabricationCheck:
    """
    The safety property: a tailored document may not assert anything the
    master doesn't support.
    """

    def test_faithful_rewrite_passes(self):
        variant = (
            "Alex Rivera\n"
            "Backend engineer with 8 years building distributed systems.\n"
            "- Designed an event pipeline processing 12M events per day"
        )

        assert verify_no_fabrication(MASTER_RESUME, variant) == []

    def test_invented_metric_is_flagged(self):
        """The classic LLM resume failure: a fabricated percentage."""
        variant = "- Reduced infrastructure costs by 60% at Northwind Systems"

        flags = verify_no_fabrication(MASTER_RESUME, variant)

        assert any("60%" in f for f in flags)

    def test_invented_employer_is_flagged(self):
        variant = "Senior Engineer, Globex Industries (2016 - 2018)"

        flags = verify_no_fabrication(MASTER_RESUME, variant)

        assert any("Globex Industries" in f for f in flags)

    def test_invented_degree_is_flagged(self):
        variant = "MSc Machine Learning, Stanford University"

        flags = verify_no_fabrication(MASTER_RESUME, variant)

        assert any("Stanford University" in f for f in flags)

    def test_altered_contact_details_flagged(self):
        variant = "Alex Rivera\nalex@totally-different.com"

        flags = verify_no_fabrication(MASTER_RESUME, variant)

        assert any("totally-different.com" in f for f in flags)

    def test_job_terms_are_allowed(self):
        """A cover letter may name the company and role it addresses."""
        variant = "I am applying for the Senior Backend Engineer role at Acme Robotics."

        flags = verify_no_fabrication(
            MASTER_RESUME, variant,
            allowed_terms=["Senior Backend Engineer", "Acme Robotics"],
        )

        assert flags == []

    def test_letter_boilerplate_is_not_flagged(self):
        variant = "Dear Hiring Manager,\n\nThank you for your consideration.\n\nSincerely,"

        assert verify_no_fabrication(MASTER_RESUME, variant) == []

    def test_numbers_already_in_master_pass(self):
        variant = "Cut query latency by 35% and raised coverage to 82%."

        assert verify_no_fabrication(MASTER_RESUME, variant) == []


# ============================================================================
# Tailoring
# ============================================================================

@pytest.mark.asyncio
class TestTailoring:
    """
    The deterministic path runs whenever no LLM is reachable — which is the
    default in tests and on any machine without `ollama serve` running.
    """

    @pytest.fixture
    def offline_service(self) -> TailoringService:
        """A service with every LLM backend disabled."""
        service = TailoringService()
        service.use_ollama = False
        service.use_anthropic = False
        return service

    async def test_deterministic_resume_is_never_fabricated(self, offline_service, job):
        result = await offline_service.tailor(MASTER_RESUME, job, DocumentType.RESUME)

        assert result.generator == "deterministic"
        assert result.fabrication_flags == []
        assert result.is_verified

    async def test_deterministic_resume_preserves_all_content(self, offline_service, job):
        """Reordering must not silently drop the candidate's experience."""
        result = await offline_service.tailor(MASTER_RESUME, job, DocumentType.RESUME)

        for expected in ("Northwind Systems", "Lumen Data", "University of Washington",
                         "alex.rivera@example.com"):
            assert expected in result.content_text

    async def test_deterministic_resume_promotes_relevant_bullets(self, offline_service, job):
        """The PostgreSQL/Kafka bullets should outrank the mentoring bullet."""
        result = await offline_service.tailor(MASTER_RESUME, job, DocumentType.RESUME)

        lines = result.content_text.split("\n")
        postgres_line = next(i for i, l in enumerate(lines) if "PostgreSQL" in l and l.strip().startswith("-"))
        mentor_line = next(i for i, l in enumerate(lines) if "Mentored" in l)

        assert postgres_line < mentor_line

    async def test_deterministic_cover_letter_names_the_job(self, offline_service, job):
        result = await offline_service.tailor(
            MASTER_RESUME, job, DocumentType.COVER_LETTER)

        assert "Senior Backend Engineer" in result.content_text
        assert "Acme Robotics" in result.content_text
        assert result.fabrication_flags == []

    async def test_deterministic_cover_letter_has_no_placeholders(self, offline_service, job):
        result = await offline_service.tailor(
            MASTER_RESUME, job, DocumentType.COVER_LETTER)

        assert "[Your Name]" not in result.content_text
        assert "Alex Rivera" in result.content_text

    async def test_notes_explain_what_happened(self, offline_service, job):
        result = await offline_service.tailor(MASTER_RESUME, job, DocumentType.RESUME)

        assert result.notes
        assert any("no llm" in n.lower() for n in result.notes)

    async def test_llm_output_is_verified_not_trusted(self, offline_service, job):
        """A model that invents a metric must be caught, not passed through."""
        async def fabricating_generator(master_text, job_record, doc_type):
            from job_agent.services.tailoring import TailoringResult
            return TailoringResult(
                content_text="- Increased revenue by 250% at Initech Corporation",
                generator="llm:test:fake",
            )

        offline_service._generate = fabricating_generator

        result = await offline_service.tailor(MASTER_RESUME, job, DocumentType.RESUME)

        assert not result.is_verified
        assert any("250%" in f for f in result.fabrication_flags)
        assert any("Initech Corporation" in f for f in result.fabrication_flags)

    async def test_unreachable_ollama_falls_back(self, job):
        """`ollama serve` being down must not break document generation."""
        service = TailoringService()
        service.use_anthropic = False
        service.use_ollama = True
        service.ollama_url = "http://127.0.0.1:9"  # nothing listens here

        result = await service.tailor(MASTER_RESUME, job, DocumentType.RESUME)

        assert result.generator == "deterministic"
        assert result.content_text

    async def test_conversational_wrapper_is_stripped(self):
        service = TailoringService()

        assert service._clean("Here is the tailored resume:\nAlex Rivera") == "Alex Rivera"
        assert service._clean("```\nAlex Rivera\n```") == "Alex Rivera"

    async def test_echoed_prompt_scaffolding_is_stripped(self):
        """
        Small models replay the prompt instead of answering it — observed with
        llama3.2:3b returning the job description and section markers verbatim.
        """
        service = TailoringService()

        echoed = (
            "--- The job being applied to ---\n"
            "Title: Senior Backend Engineer\n"
            "Description:\n"
            "Kafka and PostgreSQL work.\n"
            "--- End of inputs ---\n"
            "Alex Rivera\n"
            "Backend engineer with 8 years of experience."
        )

        cleaned = service._clean(echoed)

        assert cleaned.startswith("Alex Rivera")
        assert "Senior Backend Engineer" not in cleaned
        assert "---" not in cleaned

    async def test_echoed_output_is_rejected_as_unusable(self, job):
        service = TailoringService()

        problem = service._is_usable(
            "--- The candidate's master document ---\n" + MASTER_RESUME,
            MASTER_RESUME, job,
        )

        assert problem is not None
        assert "scaffolding" in problem

    async def test_truncated_output_is_rejected(self, job):
        service = TailoringService()

        assert service._is_usable("Alex Rivera", MASTER_RESUME, job) is not None

    async def test_output_missing_the_candidate_name_is_rejected(self, job):
        """A "resume" without the candidate's name is not their resume."""
        service = TailoringService()

        impostor = "Jordan Blake\n" + ("Backend engineer with broad experience. " * 12)

        problem = service._is_usable(impostor, MASTER_RESUME, job)

        assert problem is not None
        assert "candidate's name" in problem

    async def test_faithful_output_passes_the_usability_check(self, job):
        service = TailoringService()

        assert service._is_usable(MASTER_RESUME, MASTER_RESUME, job) is None

    async def test_output_that_drops_the_contact_header_is_rejected(self, job):
        """
        A rewrite kept the name and dropped the line under it.

        The result still parsed as a resume, rendered to a plausible PDF and
        was attached to an application — with no way for the employer to reply.
        """
        service = TailoringService()

        headerless = MASTER_RESUME.replace(
            "alex.rivera@example.com | +1 555 0100 | San Francisco, CA",
            "linkedin.com/in/alex-rivera",
        )

        problem = service._is_usable(headerless, MASTER_RESUME, job)

        assert problem is not None
        assert "email address" in problem

    async def test_a_rewrite_that_deletes_content_is_rejected(self, job):
        """
        Tailoring reorders and rewords. It does not shorten.

        Asked to tailor a resume, a small model returns a shorter one: it keeps
        what looked relevant to the posting and quietly drops the skills list,
        the second role and the languages. What comes back is not a focused
        version of the candidate's career, it is a thinner one — and the
        candidate is the last person able to see that, because they know what
        the document was meant to say.
        """
        service = TailoringService()

        halved = "\n".join(MASTER_RESUME.split("\n")[:12])

        problem = service._is_usable(halved, MASTER_RESUME, job)

        assert problem is not None
        assert "dropped" in problem

    async def test_rewording_is_not_deleting(self, job):
        """The whole point of tailoring has to survive the check."""
        service = TailoringService()

        reworded = (
            MASTER_RESUME
            .replace(
                "Led migration of the billing platform to PostgreSQL, cutting query latency by 35%",
                "Owned the billing platform's migration onto PostgreSQL, cutting query latency 35%",
            )
            .replace(
                "Designed an event pipeline processing 12M events per day",
                "Built an event-driven pipeline handling 12M events daily",
            )
            .replace(
                "Backend engineer with 8 years building distributed systems.",
                "Distributed-systems engineer, 8 years, focused on event-driven backends.",
            )
        )

        assert service._is_usable(reworded, MASTER_RESUME, job) is None

    @pytest.mark.parametrize("removed", [
        "Backend Engineer, Lumen Data (2018 - 2021)",
        "BSc Computer Science, University of Washington (2018)",
        "Python, PostgreSQL, AWS, Kafka, Docker, Kubernetes",
        "- Mentored 4 engineers across two teams",
    ])
    async def test_each_kind_of_dropped_line_is_caught(self, job, removed):
        """A role, a degree, a skills list and a single bullet all count."""
        service = TailoringService()

        without = MASTER_RESUME.replace(removed, "")

        assert service._dropped_content(without, MASTER_RESUME)

    async def test_a_cover_letter_may_be_shorter_than_the_resume(self, job):
        """A letter is not a resume; it is not held to reproducing one."""
        from job_agent.models.database import DocumentType

        service = TailoringService()

        letter = (
            "Alex Rivera\n\nDear Hiring Manager,\n\n"
            + ("I am writing to apply for the Senior Backend Engineer role, "
               "where my work on PostgreSQL and Kafka would carry over. " * 4)
            + "\n\nSincerely,\nAlex Rivera"
        )

        assert service._is_usable(
            letter, MASTER_RESUME, job, DocumentType.COVER_LETTER
        ) is None

    async def test_a_cover_letter_is_not_held_to_the_resume_header_rule(self, job):
        """A letter that recited the candidate's phone number would be worse."""
        from job_agent.models.database import DocumentType

        service = TailoringService()

        letter = (
            "Alex Rivera\n\nDear Hiring Manager,\n\n"
            + ("I am writing to apply for the Senior Backend Engineer role. " * 6)
            + "\n\nSincerely,\nAlex Rivera"
        )

        assert service._is_usable(
            letter, MASTER_RESUME, job, DocumentType.COVER_LETTER
        ) is None

    async def test_unusable_llm_output_falls_back_to_deterministic(self, job, monkeypatch):
        """Garbage from a model must never become the user's document."""
        service = TailoringService()
        service.use_ollama = True
        service.use_anthropic = False

        async def echoing_model(master_text, job_record, doc_type):
            from job_agent.services.tailoring import TailoringResult
            return TailoringResult(content_text="hi", generator="llm:ollama:tiny")

        # Simulate the model returning something unusable
        async def fake_ollama(master_text, job_record, doc_type):
            result = await echoing_model(master_text, job_record, doc_type)
            if service._is_usable(result.content_text, master_text, job_record):
                return None
            return result

        monkeypatch.setattr(service, "_tailor_with_ollama", fake_ollama)

        result = await service.tailor(MASTER_RESUME, job, DocumentType.RESUME)

        assert result.generator == "deterministic"
        assert "Northwind Systems" in result.content_text

    async def test_job_keywords_exclude_boilerplate(self, job):
        keywords = TailoringService._job_keywords(job)

        assert "postgresql" in keywords
        assert "kafka" in keywords
        assert "the" not in keywords
        assert "experience" not in keywords


# ============================================================================
# Builder: upload → tailor → render → store
# ============================================================================

class TestMasterUpload:
    """Acceptance: a master is uploaded, parsed, and stored."""

    def test_upload_stores_and_parses(self, builder, resume_docx, session):
        master = builder.upload_master(resume_docx, DocumentType.RESUME, name="My Resume")

        assert master.id is not None
        assert master.name == "My Resume"
        assert master.source_format == DocumentFormat.DOCX
        assert "Northwind Systems" in master.content_text
        assert "experience" in master.sections
        assert master.is_active

    def test_original_is_copied_into_managed_storage(self, builder, resume_txt, docs_dir):
        master = builder.upload_master(resume_txt, DocumentType.RESUME)

        stored = Path(master.source_path)
        assert stored.exists()
        assert docs_dir in stored.parents
        assert stored.read_text() == resume_txt.read_text()

    def test_upload_survives_deletion_of_the_original(self, builder, resume_txt):
        master = builder.upload_master(resume_txt, DocumentType.RESUME)
        resume_txt.unlink()

        assert Path(master.source_path).exists()

    def test_uploading_again_deactivates_the_previous_master(self, builder, resume_txt, tmp_path):
        first = builder.upload_master(resume_txt, DocumentType.RESUME, name="v1")

        second_path = tmp_path / "v2.txt"
        second_path.write_text(MASTER_RESUME)
        second = builder.upload_master(second_path, DocumentType.RESUME, name="v2")

        builder.db_session.refresh(first)

        assert not first.is_active
        assert second.is_active
        assert builder.get_active_master(DocumentType.RESUME).id == second.id

    def test_resume_and_cover_letter_are_independent(self, builder, resume_txt, tmp_path):
        resume = builder.upload_master(resume_txt, DocumentType.RESUME)

        letter_path = tmp_path / "letter.txt"
        letter_path.write_text(MASTER_COVER_LETTER)
        letter = builder.upload_master(letter_path, DocumentType.COVER_LETTER)

        builder.db_session.refresh(resume)

        assert resume.is_active  # uploading a letter must not deactivate the resume
        assert letter.is_active

    def test_oversized_upload_rejected(self, builder, tmp_path, monkeypatch):
        from job_agent.config import settings

        monkeypatch.setattr(settings, "document_max_upload_mb", 0)
        path = tmp_path / "big.txt"
        path.write_text(MASTER_RESUME)

        with pytest.raises(ValueError, match="over the"):
            builder.upload_master(path, DocumentType.RESUME)

    def test_upload_is_audited(self, builder, resume_txt, session):
        builder.upload_master(resume_txt, DocumentType.RESUME)

        entry = session.query(AuditLog).filter(
            AuditLog.action == AuditAction.DOCUMENT_UPLOADED).one()

        assert "resume" in entry.detail


@pytest.mark.asyncio
class TestVariantGeneration:
    """Acceptance: a tailored PDF variant is produced and linked."""

    async def test_generates_a_linked_pdf(self, builder, resume_txt, job, session):
        builder.upload_master(resume_txt, DocumentType.RESUME)

        version = await builder.build_variant(job, DocumentType.RESUME)

        assert version.id is not None
        assert version.job_id == job.id
        assert version.pdf_path is not None

        pdf = Path(version.pdf_path)
        assert pdf.exists()
        assert pdf.read_bytes().startswith(b"%PDF")

    async def test_variant_is_verified_against_the_master(self, builder, resume_txt, job):
        builder.upload_master(resume_txt, DocumentType.RESUME)

        version = await builder.build_variant(job, DocumentType.RESUME)

        assert version.fabrication_flags == []
        assert version.is_verified

    async def test_pdf_is_readable_and_contains_the_content(self, builder, resume_txt, job):
        builder.upload_master(resume_txt, DocumentType.RESUME)

        version = await builder.build_variant(job, DocumentType.RESUME)
        text, _sections, _fmt = parse_document(Path(version.pdf_path))

        assert "Alex Rivera" in text
        assert "Northwind" in text

    async def test_generating_without_a_master_is_a_clear_error(self, builder, job):
        with pytest.raises(ValueError, match="No active master"):
            await builder.build_variant(job, DocumentType.RESUME)

    async def test_package_includes_cover_letter_when_available(
        self, builder, resume_txt, tmp_path, job
    ):
        builder.upload_master(resume_txt, DocumentType.RESUME)

        letter_path = tmp_path / "letter.txt"
        letter_path.write_text(MASTER_COVER_LETTER)
        builder.upload_master(letter_path, DocumentType.COVER_LETTER)

        versions = await builder.build_application_package(job)

        assert [v.doc_type for v in versions] == [
            DocumentType.RESUME, DocumentType.COVER_LETTER]
        assert all(Path(v.pdf_path).exists() for v in versions)

    async def test_package_without_cover_letter_master_writes_one_from_the_resume(
        self, builder, resume_txt, job
    ):
        # A form offering a cover-letter field used to be sent it blank
        # whenever the user had uploaded no master letter. The resume is
        # already the authority on the candidate's facts, so the letter is
        # written from it — and says so in its notes.
        builder.upload_master(resume_txt, DocumentType.RESUME)

        versions = await builder.build_application_package(job)

        assert [v.doc_type for v in versions] == [
            DocumentType.RESUME, DocumentType.COVER_LETTER]

        letter = versions[1]
        assert any("master resume" in note for note in letter.tailoring_notes)

    async def test_multiple_jobs_get_separate_versions(self, builder, resume_txt, job, session):
        builder.upload_master(resume_txt, DocumentType.RESUME)

        other = Job(
            platform="generic_ats", external_id="REQ-2", title="Platform Engineer",
            company="Globex", location="Remote", description="Kubernetes and Docker",
            apply_method="web_form", dedup_hash="hash-2",
        )
        session.add(other)
        session.commit()
        session.refresh(other)

        first = await builder.build_variant(job, DocumentType.RESUME)
        second = await builder.build_variant(other, DocumentType.RESUME)

        assert first.id != second.id
        assert first.pdf_path != second.pdf_path
        assert session.query(DocumentVersion).count() == 2

    async def test_generation_is_audited(self, builder, resume_txt, job, session):
        builder.upload_master(resume_txt, DocumentType.RESUME)
        await builder.build_variant(job, DocumentType.RESUME)

        entry = session.query(AuditLog).filter(
            AuditLog.action == AuditAction.DOCUMENT_TAILORED).one()

        assert "Acme Robotics" in entry.detail
        assert entry.detail_json["verified"] is True

    async def test_flagged_variant_is_audited_as_paused(
        self, builder, resume_txt, job, session, monkeypatch
    ):
        """A variant with unsupported claims must be recorded, not silently used."""
        builder.upload_master(resume_txt, DocumentType.RESUME)

        async def fabricating(master_text, job_record, doc_type):
            from job_agent.services.tailoring import TailoringResult
            return TailoringResult(
                content_text="Alex Rivera\n- Grew revenue 400% at Initech Corporation",
                generator="llm:test:fake",
            )

        # builder.tailoring is the process-wide singleton — patch it through
        # monkeypatch so the stub is torn down before the next test runs.
        monkeypatch.setattr(builder.tailoring, "_generate", fabricating)

        version = await builder.build_variant(job, DocumentType.RESUME)

        assert not version.is_verified

        flagged = session.query(AuditLog).filter(
            AuditLog.action == AuditAction.DOCUMENT_FLAGGED).one()
        assert flagged.result == "paused"


if __name__ == "__main__":
    pytest.main([__file__, "-v"])

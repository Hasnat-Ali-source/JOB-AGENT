"""
Master Document Parser (Phase 4).

Turns an uploaded resume or cover letter into plain text plus detected
sections, so the tailoring step has structured content to work with.

Supported inputs:
- .docx  — python-docx (paragraphs + tables)
- .pdf   — pypdf text extraction
- .txt / .md — read directly

Section detection is heuristic: resumes have no standard format, so headings
are matched against the common labels and anything before the first heading is
treated as the header block (name, contact details).
"""

import logging
import re
from pathlib import Path
from typing import Dict, List, Optional, Tuple

from job_agent.models.database import DocumentFormat

logger = logging.getLogger(__name__)

# Canonical section -> heading spellings seen in the wild
SECTION_ALIASES: Dict[str, List[str]] = {
    # "PROFILE SUMMARY" is the heading on the resume this project was built
    # against and it was not in this list, so the summary section could not be
    # located at all — which silently disabled the one rewrite with the most
    # to gain. Heading vocabulary is cheap; missing one is not.
    "summary": [
        "summary", "profile", "objective", "about", "professional summary",
        "profile summary", "career summary", "personal summary",
        "professional profile", "career profile", "executive summary",
        "career objective", "personal statement", "overview",
    ],
    "experience": [
        "experience", "work experience", "employment", "professional experience",
        "work history", "career history", "employment history",
        "relevant experience", "professional background",
    ],
    "education": [
        "education", "academic background", "qualifications",
        "education and training", "academic qualifications",
    ],
    "skills": [
        "skills", "technical skills", "core skills", "key skills",
        "core competencies", "competencies", "technologies", "skills summary",
    ],
    "projects": ["projects", "selected projects", "personal projects"],
    "certifications": ["certifications", "certificates", "licenses"],
    "awards": ["awards", "honors", "achievements"],
    "publications": ["publications", "papers"],
    "languages": ["languages"],
    "interests": ["interests", "hobbies"],
}

# A heading is short, mostly letters, and often upper case or title case
_MAX_HEADING_WORDS = 5

EMAIL_PATTERN = re.compile(r"[\w.+-]+@[\w-]+\.[\w.]+")

# Deliberately loose about formatting — international spacing, brackets and
# dashes all appear in real headers — and deliberately strict about length,
# because "2017 - 2021" is otherwise a phone number as far as a regex is
# concerned, and a resume that lost its header would look like it kept one.
_PHONE_PATTERN = re.compile(r"\+?\d[\d\s().-]{6,}\d")


def has_email_address(text: str) -> bool:
    """
    Whether an employer reading this document could email the candidate.

    Args:
        text: Document text

    Returns:
        True if an email address appears anywhere
    """
    return bool(EMAIL_PATTERN.search(text or ""))


def has_phone_number(text: str) -> bool:
    """
    Whether a phone number appears in this document.

    A run of digits counts only if it carries a country prefix or is long
    enough to be a number rather than a date range.

    Args:
        text: Document text

    Returns:
        True if a phone number appears anywhere
    """
    for match in _PHONE_PATTERN.finditer(text or ""):
        found = match.group(0)

        if found.startswith("+") or len(re.sub(r"\D", "", found)) >= 9:
            return True

    return False


class DocumentParser:
    """Extracts text and sections from master documents."""

    @staticmethod
    def detect_format(path: Path) -> DocumentFormat:
        """
        Determine the document format from the file suffix.

        Args:
            path: Path to the document

        Returns:
            DocumentFormat

        Raises:
            ValueError: If the extension isn't supported
        """
        suffix = path.suffix.lower().lstrip(".")

        mapping = {
            "docx": DocumentFormat.DOCX,
            "pdf": DocumentFormat.PDF,
            "txt": DocumentFormat.TXT,
            "md": DocumentFormat.MARKDOWN,
            "markdown": DocumentFormat.MARKDOWN,
        }

        if suffix not in mapping:
            raise ValueError(
                f"Unsupported document format '.{suffix}' — "
                f"upload a .docx, .pdf, .txt or .md file"
            )

        return mapping[suffix]

    @staticmethod
    def parse(path: Path) -> Tuple[str, Dict[str, str], DocumentFormat]:
        """
        Parse a document into text and sections.

        Args:
            path: Path to the document

        Returns:
            (full_text, sections, format)

        Raises:
            FileNotFoundError: If the file doesn't exist
            ValueError: If the format is unsupported or the file has no text
        """
        if not path.exists():
            raise FileNotFoundError(f"Document not found: {path}")

        doc_format = DocumentParser.detect_format(path)

        if doc_format == DocumentFormat.DOCX:
            text = DocumentParser._parse_docx(path)
        elif doc_format == DocumentFormat.PDF:
            text = DocumentParser._parse_pdf(path)
        else:
            text = path.read_text(encoding="utf-8", errors="replace")

        text = DocumentParser._normalize(text)

        if not text.strip():
            raise ValueError(
                f"No text could be extracted from {path.name}. "
                f"If this is a scanned PDF, it needs OCR first."
            )

        sections = DocumentParser.split_sections(text)

        logger.info(
            f"Parsed {path.name}: {len(text)} chars, "
            f"sections: {', '.join(sections) or 'none detected'}"
        )

        return text, sections, doc_format

    @staticmethod
    def _parse_docx(path: Path) -> str:
        """Extract paragraphs and table cells from a .docx file."""
        from docx import Document

        document = Document(str(path))
        lines: List[str] = []

        for paragraph in document.paragraphs:
            lines.append(paragraph.text)

        # Many resume templates lay content out in tables
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    lines.append(" | ".join(cells))

        return "\n".join(lines)

    @staticmethod
    def _parse_pdf(path: Path) -> str:
        """
        Extract text from a .pdf file, in the order a reader sees it.

        The default extraction returns text in the order the PDF's content
        stream happens to draw it, which for a designed resume is not the order
        it is read in. On a real one it put the candidate's contact line, the
        first sentence of their summary, and one of their job titles at the
        very end of the document — so the summary began mid-sentence, a list of
        bullets sat under a heading with no employer above it, and the email
        address appeared after the hobbies.

        Nothing downstream could recover from that: the tailored resume, the
        fabrication check and the rendered PDF all inherited the scrambled
        order. Layout mode uses each fragment's position on the page instead,
        which is the order the document was written in.

        Args:
            path: The PDF

        Returns:
            The document's text
        """
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        pages = []

        for page in reader.pages:
            text = ""

            try:
                text = page.extract_text(extraction_mode="layout") or ""
            except Exception as e:
                logger.info(
                    f"Layout extraction failed for a page of {path.name} "
                    f"({type(e).__name__}) — falling back to stream order"
                )

            # Layout mode returns nothing on a PDF whose text it cannot place;
            # stream order is wrong-but-present, which beats an empty resume.
            if not text.strip():
                try:
                    text = page.extract_text() or ""
                except Exception as e:
                    logger.warning(f"Could not extract a page from {path.name}: {e}")
                    continue

            pages.append(text)

        return "\n".join(pages)

    @staticmethod
    def _normalize(text: str) -> str:
        """Collapse excess whitespace while preserving line structure."""
        text = text.replace("\r\n", "\n").replace("\r", "\n")

        # Designed resumes set their contact icons — envelope, phone, pin — in
        # an icon font. Extracted, those glyphs come out as private-use
        # codepoints or as the substitute square the viewer drew instead, and
        # they are then rendered into the tailored PDF as "■ hello@example.com".
        # The icon carried no information the text does not.
        text = re.sub(
            "[-"       # private use area — where icon fonts live
            "■-◿"        # geometric shapes, the usual substitute
            "�]",             # the replacement character itself
            "",
            text,
        )

        text = re.sub(r"[ \t]+", " ", text)
        text = re.sub(r"\n{3,}", "\n\n", text)

        lines = [line.strip() for line in text.split("\n")]

        return DocumentParser._rejoin_wrapped_lines(lines).strip()

    @staticmethod
    def _rejoin_wrapped_lines(lines: List[str]) -> str:
        """
        Put sentences broken by the page's width back together.

        A PDF has no paragraphs, only lines placed on a page, so a sentence
        that ran past the margin arrives as two. Rendered back into a new
        document at a different width, each fragment becomes its own paragraph:
        a four-line summary printed as four stubs, and a bullet whose last word
        ("scores.") sits alone underneath it. It looks like a document that was
        assembled carelessly, which is the one impression a resume cannot
        afford.

        A break is a wrap — not a real line ending — when the line stops
        without terminal punctuation and the next one continues in lower case.
        That is conservative on purpose: an address block, a skills list and a
        run of job titles all start their lines with a capital or a digit, so
        none of them is joined.

        Args:
            lines: Stripped lines, in reading order

        Returns:
            The text with wrapped lines rejoined
        """
        joined: List[str] = []

        for line in lines:
            if not joined or not line:
                joined.append(line)
                continue

            previous = joined[-1]

            if DocumentParser._is_wrapped(previous, line):
                joined[-1] = f"{previous} {line}"
                continue

            joined.append(line)

        return "\n".join(joined)

    @staticmethod
    def _is_wrapped(previous: str, line: str) -> bool:
        """
        Whether `line` is the rest of `previous` rather than a new one.

        Args:
            previous: The line already collected
            line: The line being considered

        Returns:
            True if the two are one line broken by the page width
        """
        if not previous or not line:
            return False

        # A bullet, a heading or a new entry starts something; it never
        # continues the line above.
        if line[0] in "-•*·–—" or DocumentParser.is_heading(line):
            return False

        # A continuation carries on in lower case. A capital or a digit starts
        # a new name, title, date or list item.
        if not line[0].islower():
            return False

        # Links and addresses are lower case and are nobody's second half —
        # "linkedin.com/in/…" on its own line is a contact detail, not the rest
        # of the phone number above it.
        if re.match(r"(https?://|www\.|[\w.+-]+@|[\w-]+\.(com|io|net|org|dev)/)", line):
            return False

        # The line above has to look unfinished. Terminal punctuation means it
        # said what it had to say, whatever comes next.
        if previous.endswith((".", "!", "?", ":", ";")) or DocumentParser.is_heading(previous):
            return False

        return True

    @staticmethod
    def is_heading(line: str) -> Optional[str]:
        """
        Return the canonical section name if this line is a section heading.

        Args:
            line: A single line of text

        Returns:
            Canonical section name, or None
        """
        stripped = line.strip().strip(":").strip()

        if not stripped or len(stripped.split()) > _MAX_HEADING_WORDS:
            return None

        # Ignore bullet lines and anything with sentence punctuation
        if stripped[0] in "-•*·" or stripped.endswith((".", ",", ";")):
            return None

        lowered = stripped.lower()

        for canonical, aliases in SECTION_ALIASES.items():
            if lowered in aliases:
                return canonical

        return None

    @staticmethod
    def split_sections(text: str) -> Dict[str, str]:
        """
        Split document text into named sections.

        Content before the first recognized heading becomes "header" — on a
        resume that's the name and contact block.

        Args:
            text: Normalized document text

        Returns:
            {section_name: section_text}
        """
        sections: Dict[str, List[str]] = {}
        current = "header"

        for line in text.split("\n"):
            heading = DocumentParser.is_heading(line)

            if heading:
                current = heading
                sections.setdefault(current, [])
                continue

            sections.setdefault(current, []).append(line)

        return {
            name: "\n".join(lines).strip()
            for name, lines in sections.items()
            if "\n".join(lines).strip()
        }


def parse_document(path: Path) -> Tuple[str, Dict[str, str], DocumentFormat]:
    """
    Convenience wrapper around DocumentParser.parse().

    Args:
        path: Path to the document

    Returns:
        (full_text, sections, format)
    """
    return DocumentParser.parse(path)
